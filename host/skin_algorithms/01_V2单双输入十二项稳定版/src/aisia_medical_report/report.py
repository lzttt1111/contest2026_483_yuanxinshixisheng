from __future__ import annotations

from pathlib import Path
from typing import Any, Iterable

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.image.image import Image as DocxImage
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


CYAN = "41C6DE"
DARK = "1F2937"
LIGHT = "EAF8FB"
WARNING = "FFF4E5"
GRAY = "F3F4F6"
FORMAL_RESULT_IMAGE_WIDTH = Cm(13.72)

_FORMAL_RESULT_CAPTIONS = {
    "01": "可见毛孔分布",
    "02": "油脂分泌倾向分布",
    "03": "综合色素分布",
    "04": "弥漫性泛红分布",
    "05": "血管样结构分布",
    "06": "痤疮样特征分布",
    "07": "干燥性细纹分布",
    "08": "稳定性皱纹分布",
    "09": "结构性沟纹分布",
    "10": "表面不规则分布",
    "11": "面部轮廓几何测量",
}


def render_docx(payload: dict[str, Any], template_path: str | Path, output_path: str | Path) -> Path:
    """兼容旧入口：旧 ``docx`` 键仍指向医生详细版。"""
    return render_doctor_docx(payload, template_path, output_path)


def _new_document(template_path: str | Path) -> DocumentType:
    template = Path(template_path).expanduser().resolve()
    if not template.is_file():
        raise FileNotFoundError(f"报告模板不存在: {template}")
    doc = Document(str(template))
    _clear_document_body(doc)
    _clear_headers_and_footers(doc)
    properties = doc.core_properties
    properties.title = "AISIA面部皮肤检测报告"
    properties.subject = "面部多指标检测结果"
    properties.keywords = "AISIA,面部检测,皮肤指标"
    properties.comments = ""
    _configure_document(doc)
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")
    return doc


def render_user_docx(payload: dict[str, Any], template_path: str | Path, output_path: str | Path) -> Path:
    target = Path(output_path).expanduser().resolve()
    target.parent.mkdir(parents=True, exist_ok=True)
    doc = _new_document(template_path)
    _add_v011_cover(doc, payload, "用户精简报告")
    _add_user_overview(doc, payload, page_break_before=True)
    _add_clickable_directory(doc, payload)
    for index, module in enumerate(payload.get("检测模块", []), 1):
        _add_user_module(doc, module, index, page_break_before=True)
    _finalize_tables(doc)
    doc.save(str(target))
    return target


def render_doctor_docx(payload: dict[str, Any], template_path: str | Path, output_path: str | Path) -> Path:
    target = Path(output_path).expanduser().resolve()
    target.parent.mkdir(parents=True, exist_ok=True)
    doc = _new_document(template_path)
    _add_v011_cover(doc, payload, "医生详细报告")
    _add_doctor_overview(doc, payload, page_break_before=True)
    _add_clickable_directory(doc, payload)
    for index, module in enumerate(payload.get("检测模块", []), 1):
        _add_doctor_module(doc, module, index, page_break_before=True)

    _finalize_tables(doc)
    doc.save(str(target))
    return target


def _clean_visible(value: Any) -> str:
    text = "" if value is None else str(value)
    replacements = {
        "待医学标定": "—", "不生成": "—", "不可评估": "—",
        "本次未检测/待扩展": "检测完成", "本次未检测": "检测完成", "部分支持": "检测完成",
        "candidate/uncalibrated": "", "legacy_shadow_score": "",
        "AI参考结果": "检测结果", "AI参考": "检测结果",
    }
    for source, target in replacements.items():
        text = text.replace(source, target)
    return text


def _add_v011_cover(doc: DocumentType, payload: dict[str, Any], subtitle: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(75)
    run = p.add_run("AISIA")
    run.bold = True
    run.font.size = Pt(36)
    run.font.color.rgb = RGBColor(65, 198, 222)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    report_title = "面部皮肤检测报告" if subtitle == "用户精简报告" else "面部多指标检测医生详细报告"
    run = title.add_run(f"AISIA｜{report_title}")
    run.bold = True
    run.font.size = Pt(26)
    info = payload.get("报告信息", {})
    subject = payload.get("受检者信息", {})
    _key_value_table(doc, [
        ("报告编号", info.get("报告编号", "")),
        ("受检者编号", subject.get("姓名或编号", "匿名受检者")),
        ("采集日期", subject.get("采集日期", info.get("报告日期", ""))),
        ("报告日期", info.get("报告日期", subject.get("采集日期", ""))),
    ])
    _add_cover_capture_image(doc, payload)


def _add_cover_capture_image(doc: DocumentType, payload: dict[str, Any]) -> None:
    """Add the optional, unannotated standard capture to the cover page."""

    artifact = payload.get("标准采集图像") or {}
    path = Path(str(artifact.get("path", "")))
    if not path.is_file():
        return
    image = DocxImage.from_file(str(path))
    if not image.px_width or not image.px_height:
        raise ValueError(f"unable to read standard capture dimensions: {path}")
    aspect = float(image.px_width) / float(image.px_height)
    maximum_width = 4.5
    maximum_height = 4.6
    width = min(maximum_width, maximum_height * aspect)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.add_run().add_picture(str(path), width=Inches(width))
    caption = doc.add_paragraph(
        _clean_visible(artifact.get("caption") or "标准白光正面图")
    )
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(0)


def _formal_modules(payload: dict[str, Any]) -> list[dict[str, Any]]:
    return sorted(
        [module for module in payload.get("检测模块", []) if isinstance(module.get("综合得分"), (int, float))],
        key=lambda row: float(row["综合得分"]), reverse=True,
    )


def _unscored_result_text(module: dict[str, Any]) -> str:
    summary = _clean_visible(module.get("结果摘要"))
    return "未检出" if "未检出" in summary else "暂不评分"


def _add_user_overview(
    doc: DocumentType,
    payload: dict[str, Any],
    *,
    page_break_before: bool = False,
) -> None:
    _bookmark_heading(
        doc,
        "结果总览",
        "user_overview",
        level=1,
        page_break_before=page_break_before,
    )
    table = doc.add_table(rows=1, cols=3)
    _header_row(table.rows[0].cells, ["检测项目", "结果", "重点表现"])
    for module in payload.get("检测模块", []):
        score = module.get("综合得分")
        result = f"{float(score):.1f}分·{module.get('程度等级')}" if isinstance(score, (int, float)) else _unscored_result_text(module)
        status = module.get("用户评估状态", module.get("评估状态", "检测完成"))
        values = [module.get("模块名称"), f"{status}｜{result}", _clean_visible(module.get("结果摘要"))]
        for cell, value in zip(table.add_row().cells, values):
            cell.text = _clean_visible(value)
            _set_cell_font(cell, size=8.5)
    doc.add_heading("重点关注", level=2)
    for module in _formal_modules(payload)[:3]:
        doc.add_paragraph(
            f"{module.get('模块名称')}：{float(module.get('综合得分')):.1f}分，"
            f"{module.get('程度等级')}。{_clean_visible(module.get('结果摘要'))}",
            style="List Bullet",
        )


def _add_doctor_overview(
    doc: DocumentType,
    payload: dict[str, Any],
    *,
    page_break_before: bool = False,
) -> None:
    _bookmark_heading(
        doc,
        "多指标结果总览",
        "doctor_overview",
        level=1,
        page_break_before=page_break_before,
    )
    table = doc.add_table(rows=1, cols=4)
    _header_row(table.rows[0].cells, ["检测维度", "评估结果", "分数/等级", "结果摘要"])
    for module in payload.get("检测模块", []):
        score = module.get("综合得分")
        score_text = f"{float(score):.2f} / {module.get('程度等级')}" if isinstance(score, (int, float)) else _unscored_result_text(module)
        status = module.get("医生评估状态", module.get("评估状态", "检测完成"))
        values = [module.get("模块名称"), status, score_text, module.get("结果摘要")]
        for cell, value in zip(table.add_row().cells, values):
            cell.text = _clean_visible(value)
            _set_cell_font(cell, size=8)


def _add_toc_field(doc: DocumentType) -> None:
    p = doc.add_paragraph()
    run = p.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText"); instruction.set(qn("xml:space"), "preserve"); instruction.text = 'TOC \\o "1-2" \\h \\z \\u'
    separate = OxmlElement("w:fldChar"); separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t"); text.text = ""
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    for node in (begin, instruction, separate, text, end):
        run._r.append(node)


def _bookmark_heading(
    doc: DocumentType,
    text: str,
    bookmark: str,
    *,
    level: int,
    page_break_before: bool = False,
) -> None:
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.page_break_before = page_break_before
    bookmark_id = str(sum(ord(ch) for ch in bookmark) + len(doc.paragraphs) * 1000)
    start = OxmlElement("w:bookmarkStart"); start.set(qn("w:id"), bookmark_id); start.set(qn("w:name"), bookmark)
    end = OxmlElement("w:bookmarkEnd"); end.set(qn("w:id"), bookmark_id)
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def _add_hyperlink(doc: DocumentType, text: str, anchor: str) -> None:
    paragraph = doc.add_paragraph()
    hyperlink = OxmlElement("w:hyperlink"); hyperlink.set(qn("w:anchor"), anchor)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color"); color.set(qn("w:val"), "2B8AA0")
    underline = OxmlElement("w:u"); underline.set(qn("w:val"), "single")
    props.extend((color, underline)); run.append(props)
    node = OxmlElement("w:t"); node.text = text; run.append(node); hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _add_clickable_directory(doc: DocumentType, payload: dict[str, Any]) -> None:
    _bookmark_heading(doc, "目录", "report_toc", level=1)
    _add_toc_field(doc)
    for index, module in enumerate(payload.get("检测模块", []), 1):
        _add_hyperlink(doc, f"{index:02d}  {module.get('模块名称')}", f"module_{index:02d}")


def _usable_metrics(module: dict[str, Any], limit: int | None = None) -> list[dict[str, Any]]:
    rows = []
    for metric in module.get("核心指标", []):
        value = metric.get("value")
        if value in (None, "", "不可评估", "—"):
            continue
        rows.append(metric)
        if limit is not None and len(rows) >= limit:
            break
    return rows


def _result_image_caption(module: dict[str, Any], path: Path) -> str:
    module_id = str(module.get("模块编号", ""))
    if module_id == "03":
        if "可见斑点" in path.name:
            return "可见色斑结果"
        if "UV色斑" in path.name:
            return "UV色斑结果"
        if "红褐混合" in path.name:
            return "红褐混合印记"
        if "重点色斑" in path.name:
            return "重点色斑区域"
    if module_id == "10":
        return "表面不规则分布" if "不规则" in path.name else "表面纹理分布"
    if str(module.get("模块编号", "")) == "02":
        lowered = path.name.lower()
        if "油光" in path.name or "gloss" in lowered:
            return "表面油光分布"
        if "卟啉" in path.name or "荧光" in path.name or "follicular_fluorescence" in lowered:
            return "毛囊荧光分布"
    return _FORMAL_RESULT_CAPTIONS.get(
        str(module.get("模块编号", "")),
        path.name,
    )


def _user_result_image_limit(module: dict[str, Any]) -> int:
    """精简版减少解释和表格，不删除组成正式结论的核心证据图。"""
    return {
        "02": 2,
        "03": 5,
        "10": 2,
    }.get(str(module.get("模块编号", "")), 1)


def _add_result_image(doc: DocumentType, module: dict[str, Any], *, maximum: int) -> None:
    paths = [Path(value) for value in module.get("结果图", []) if Path(value).is_file()]
    if not paths:
        return
    doc.add_heading("结果图", level=2)
    selected = paths[:maximum]
    for path in selected:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=FORMAL_RESULT_IMAGE_WIDTH)
        caption_text = _result_image_caption(module, path)
        caption = doc.add_paragraph(caption_text); caption.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _format_metric_value(metric: dict[str, Any]) -> str:
    """Render only explicitly tagged report metrics; never infer a unit from value."""
    value = metric.get("value")
    display_type = metric.get("display_type", "TEXT")
    if value is None:
        return "—"
    if display_type == "PERCENTAGE" and isinstance(value, (int, float)):
        return f"{float(value) * 100:.2f}"
    if display_type == "COUNT" and isinstance(value, (int, float)):
        return str(int(round(float(value))))
    if display_type in {"NORMALIZED_0_1", "ENGINEERING_0_1", "INTENSITY_ENGINEERING"} and isinstance(value, (int, float)):
        return f"{float(value):.4f}"
    if display_type == "RATIO" and isinstance(value, (int, float)):
        return f"{float(value):.2f}"
    if display_type in {"DENSITY", "DENSITY_100K", "DENSITY_10K", "LENGTH_DENSITY_10K"} and isinstance(value, (int, float)):
        return f"{float(value):.2f}"
    if display_type == "AREA_PIXEL" and isinstance(value, (int, float)):
        number = float(value)
        return f"{number:.0f}" if number.is_integer() else f"{number:.2f}"
    if display_type == "LENGTH_PIXEL" and isinstance(value, (int, float)):
        number = float(value)
        return f"{number:.0f}" if number.is_integer() else f"{number:.1f}"
    if display_type == "DELTA_E" and isinstance(value, (int, float)):
        return f"{float(value):.2f}"
    if display_type == "SCORE" and isinstance(value, (int, float)):
        return f"{float(value):.2f}"
    return _format_value(value)


def _add_metric_table(doc: DocumentType, metrics: list[dict[str, Any]], *, doctor: bool = False) -> None:
    table = doc.add_table(rows=1, cols=4)
    _header_row(table.rows[0].cells, ["核心指标", "检测结果", "单位", "结果解释"])
    for metric in metrics:
        values = [
            metric.get("name"),
            _format_metric_value(metric),
            _clean_visible(metric.get("unit", "")) or "—",
            metric.get("explanation", ""),
        ]
        for cell, value in zip(table.add_row().cells, values):
            cell.text = _clean_visible(value)
            _set_cell_font(cell, size=8.0 if doctor else 8.5)


def _add_user_module(
    doc: DocumentType,
    module: dict[str, Any],
    index: int,
    *,
    page_break_before: bool = False,
) -> None:
    _bookmark_heading(
        doc,
        f"{index:02d}  {module.get('模块名称')}",
        f"module_{index:02d}",
        level=1,
        page_break_before=page_break_before,
    )
    score = module.get("综合得分")
    if isinstance(score, (int, float)):
        _key_value_table(doc, [("问题负担分", f"{float(score):.2f}"), ("程度等级", module.get("程度等级"))])
    doc.add_paragraph(_clean_visible(module.get("结果摘要")))
    _add_result_image(doc, module, maximum=_user_result_image_limit(module))
    metrics = _usable_metrics(module, 4)
    if metrics:
        doc.add_heading("核心指标", level=2)
        _add_metric_table(doc, metrics)
    note = _clean_visible(module.get("日常管理提示"))
    if note:
        doc.add_heading("日常建议", level=2)
        doc.add_paragraph(note)


def _add_doctor_module(
    doc: DocumentType,
    module: dict[str, Any],
    index: int,
    *,
    page_break_before: bool = False,
) -> None:
    _bookmark_heading(
        doc,
        f"{index:02d}  {module.get('模块名称')}",
        f"module_{index:02d}",
        level=1,
        page_break_before=page_break_before,
    )
    score = module.get("综合得分")
    rows: list[tuple[str, Any]] = [
        ("评估状态", module.get("医生评估状态", module.get("评估状态", "检测完成"))),
        ("数据来源", "、".join(module.get("数据来源", []))),
    ]
    if isinstance(score, (int, float)):
        rows.extend([("\u95ee\u9898\u8d1f\u62c5\u5206", "%.2f" % float(score)), ("\u7a0b\u5ea6\u7b49\u7ea7", module.get("\u7a0b\u5ea6\u7b49\u7ea7"))])
    else:
        rows.append(("检测结论", _unscored_result_text(module)))
    _key_value_table(doc, rows)
    doc.add_heading("结果摘要", level=2)
    doc.add_paragraph(_clean_visible(module.get("结果摘要")))
    has_group_images = any(group.get("images") for group in module.get("医生结果分组", []) or [])
    if not has_group_images:
        _add_result_image(doc, module, maximum=1)
    metrics = [] if module.get("医生版隐藏全面部核心指标") else _usable_metrics(module)
    if metrics:
        doc.add_heading("全面部核心指标", level=2)
        _add_metric_table(doc, metrics, doctor=True)
    for group in module.get("医生结果分组", []) or []:
        doc.add_heading(_clean_visible(group.get("title")), level=2)
        summary = _clean_visible(group.get("summary"))
        if summary:
            doc.add_paragraph(summary)
        group_images = [
            (Path(image.get("path", "")), _clean_visible(image.get("caption", "")))
            for image in group.get("images", []) or []
            if Path(image.get("path", "")).is_file()
        ]
        for image in group.get("images", []) or []:
            path = Path(image.get("path", ""))
            if not path.is_file():
                continue
            paragraph = doc.add_paragraph(); paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.add_run().add_picture(str(path), width=FORMAL_RESULT_IMAGE_WIDTH)
            caption = doc.add_paragraph(_clean_visible(image.get("caption", path.stem)))
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        group_metrics = _usable_metrics({"核心指标": group.get("metrics", [])})
        if group_metrics:
            _add_metric_table(doc, group_metrics, doctor=True)
        elif group.get("unavailable"):
            doc.add_paragraph("当前版本尚未形成可靠量化。")
    regions = module.get("分区指标", [])
    if regions:
        doc.add_heading("分区结果", level=2)
        table = doc.add_table(rows=1, cols=3)
        _header_row(table.rows[0].cells, ["面部分区", "状态", "主要指标"])
        for region in regions:
            values = [region.get("分区名称"), region.get("状态", "检测完成"), _compact_metrics(region.get("指标", {}), limit=12)]
            for cell, value in zip(table.add_row().cells, values):
                cell.text = _clean_visible(value)
                _set_cell_font(cell, size=7.5)
    left_right_text = _clean_visible(module.get("左右比较摘要"))
    if left_right_text:
        doc.add_heading("左右比较", level=2)
        doc.add_paragraph(left_right_text)
    explanation = module.get("正式评分说明")
    if isinstance(explanation, dict) and explanation.get("summary_text"):
        doc.add_heading("评分说明", level=2)
        doc.add_paragraph(_clean_visible(explanation["summary_text"]))
        drivers = explanation.get("drivers") or []
        if drivers:
            table = doc.add_table(rows=1, cols=4)
            _header_row(table.rows[0].cells, ["评分因素", "权重", "分项得分", "贡献"])
            for driver in drivers:
                values = [
                    driver.get("name"), f"{float(driver.get('weight', 0))*100:.1f}%",
                    f"{float(driver.get('score', 0)):.2f}",
                    f"{float(driver.get('weighted_contribution', 0)):.2f}",
                ]
                for cell, value in zip(table.add_row().cells, values):
                    cell.text = _clean_visible(value)
                    _set_cell_font(cell, size=7.5)


def _add_technical_appendix(doc: DocumentType, payload: dict[str, Any]) -> None:
    _bookmark_heading(doc, "附录：版本与追溯信息", "doctor_appendix", level=1)
    summary = payload.get("评分配置摘要", {})
    rows = [
        ("报告数据版本", payload.get("报告数据版本", "report_payload_v011")),
        ("评分配置版本", summary.get("scoring_profile_version", "")),
        ("评分配置SHA-256", summary.get("profile_sha256", "")),
        ("指标注册表SHA-256", summary.get("registry_sha256", "")),
        ("算法证据版本", summary.get("algorithm_version", "")),
        ("指标Schema版本", summary.get("metrics_schema_version", "")),
        ("报告编号", payload.get("报告信息", {}).get("报告编号", "")),
    ]
    _key_value_table(doc, rows)
    doc.add_heading("采集与复测要点", level=2)
    for note in (
        "使用同一设备、光源、拍摄距离和曝光参数进行纵向复测。",
        "保持面部清洁状态、表情和头部角度一致，并记录采集时间。",
        "同名指标的纵向比较应使用相同评分配置和算法证据版本。",
    ):
        doc.add_paragraph(note, style="List Bullet")


def _clear_document_body(doc: DocumentType) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def _clear_headers_and_footers(doc: DocumentType) -> None:
    """Remove all template demo text before rendering a formal report."""
    for section in doc.sections:
        for container in (section.header, section.first_page_header, section.even_page_header,
                          section.footer, section.first_page_footer, section.even_page_footer):
            for paragraph in container.paragraphs:
                paragraph.clear()
            for table in list(container.tables):
                table._element.getparent().remove(table._element)


def _configure_document(doc: DocumentType) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    for style_name, size in (("Normal", 10.5), ("Title", 28), ("Heading 1", 20), ("Heading 2", 15)):
        style = doc.styles[style_name]
        style.font.name = "微软雅黑"
        style.font.size = Pt(size)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def _add_cover(doc: DocumentType, payload: dict[str, Any]) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(80)
    run = p.add_run("AISIA")
    run.bold = True
    run.font.size = Pt(34)
    run.font.color.rgb = RGBColor(65, 198, 222)
    run.font.name = "Arial"

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("面部多指标检测\n医学量化汇总报告")
    run.bold = True
    run.font.size = Pt(27)
    run.font.color.rgb = RGBColor(31, 41, 55)

    info = payload.get("报告信息", {})
    subject = payload.get("受检者信息", {})
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    rows = [
        ("报告编号", info.get("报告编号", "")),
        ("受检者编号", subject.get("姓名或编号", subject.get("姓名", "未提供"))),
        ("采集日期", subject.get("采集日期", "未提供")),
        ("生成时间", info.get("生成时间", "")),
    ]
    if info.get("评分模式") != "v0.1.1_integrity":
        rows.extend((("综合得分", info.get("综合得分", "待医学标定")), ("程度等级", info.get("程度等级", "待医学标定"))))
    for label, val in rows:
        cells = table.add_row().cells
        cells[0].width = Cm(4)
        cells[1].width = Cm(10)
        cells[0].text = str(label)
        cells[1].text = str(val)
        _shade(cells[0], LIGHT)
        _set_cell_font(cells[0], bold=True)
        _set_cell_font(cells[1])

    doc.add_paragraph()
    warning = doc.add_table(rows=1, cols=1)
    warning.alignment = WD_TABLE_ALIGNMENT.CENTER
    warning.cell(0, 0).text = "重要说明：本报告为AI辅助面部图像量化结果，不构成医学诊断、治疗建议或疾病严重程度判定。"
    _shade(warning.cell(0, 0), WARNING)
    _set_cell_font(warning.cell(0, 0), bold=True)


def _add_report_notice(doc: DocumentType, payload: dict[str, Any]) -> None:
    doc.add_heading("报告说明", level=1)
    info = payload.get("报告信息", {})
    quality = payload.get("采集与质量信息", {})
    rows = [
        ("报告范围", f"共包含11个医学章节，其中 {payload.get('可用模块数量', 0)} 个章节具备当前算法结果，其余明确标记为本次未检测或待扩展。"),
        ("评分说明", "综合得分和程度等级尚未完成人群常模与医生标注校准，当前统一显示为“待医学标定”。"),
        ("成像边界", quality.get("成像说明", "普通白光RGB二维图像")),
        ("左右口径", quality.get("左右说明", "画面左侧/画面右侧")),
        ("报告性质", info.get("报告性质", "AI辅助图像量化报告")),
    ]
    _key_value_table(doc, rows)


def _add_overview(doc: DocumentType, payload: dict[str, Any]) -> None:
    doc.add_heading("多指标结果总览", level=1)
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _header_row(table.rows[0].cells, ["检测模块", "当前状态", "评分/等级", "汇总表现"])
    for module in payload.get("检测模块", []):
        cells = table.add_row().cells
        score = module.get("综合得分")
        grade = module.get("程度等级")
        score_text = f"{score:.2f} / {grade}" if isinstance(score, (int, float)) else "仅展示检测结果"
        values = [module.get("模块名称", ""), module.get("可用状态", ""), score_text, module.get("结果摘要", "")]
        for cell, val in zip(cells, values):
            cell.text = str(val)
            _set_cell_font(cell, size=8.5)


def _add_cross_summary(doc: DocumentType, payload: dict[str, Any]) -> None:
    doc.add_heading("汇总重点", level=1)
    for item in payload.get("总体摘要", []):
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(f"{item.get('模块', '')}：")
        run.bold = True
        p.add_run(str(item.get("重点", "")))


def _add_directory(doc: DocumentType, payload: dict[str, Any]) -> None:
    doc.add_heading("目录", level=1)
    doc.add_paragraph("以下目录保留医生模板中的11个医学章节。未支持章节仍保留并明确说明原因。")
    for module in payload.get("检测模块", []):
        doc.add_paragraph(f"{module.get('模块编号')}  {module.get('模块名称')}  —  {module.get('可用状态')}")


def _add_module(doc: DocumentType, module: dict[str, Any]) -> None:
    title = f"{module.get('模块编号')}  {module.get('模块名称')}"
    doc.add_heading(title, level=1)
    table = doc.add_table(rows=1, cols=4)
    _header_row(table.rows[0].cells, ["模块", "当前状态", "综合得分", "程度等级"])
    row = table.add_row().cells
    for cell, val in zip(row, [
        module.get("模块名称"), module.get("可用状态"),
        module.get("综合得分") if module.get("综合得分") is not None else "不生成",
        module.get("程度等级") if module.get("程度等级") is not None else "不生成",
    ]):
        cell.text = str(val)
        _set_cell_font(cell, bold=cell is row[0])

    doc.add_heading("结果摘要", level=2)
    doc.add_paragraph(str(module.get("结果摘要", "")))

    images = [Path(p) for p in module.get("结果图", []) if Path(p).is_file()]
    if images:
        doc.add_heading("检测结果图", level=2)
        for image in images[:3]:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            try:
                p.add_run().add_picture(str(image), width=Inches(5.8))
                caption = doc.add_paragraph(image.name)
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            except Exception as exc:
                doc.add_paragraph(f"图片无法加载：{image.name}（{exc}）")
    else:
        p = doc.add_paragraph("本章节无可用的同图结果图片。")
        if p.runs:
            p.runs[0].italic = True

    doc.add_heading("核心指标", level=2)
    metrics = module.get("核心指标", [])
    if metrics:
        table = doc.add_table(rows=1, cols=4)
        _header_row(table.rows[0].cells, ["核心指标", "检测结果", "单位", "说明"])
        for metric in metrics:
            cells = table.add_row().cells
            vals = [metric.get("name", ""), _format_value(metric.get("value")), metric.get("unit", ""), metric.get("note", "")]
            for cell, val in zip(cells, vals):
                cell.text = str(val)
                _set_cell_font(cell, size=8.5)
    else:
        doc.add_paragraph("本章节暂无可量化指标。")

    doc.add_heading("分区结果", level=2)
    regions = module.get("分区指标", [])
    if regions:
        table = doc.add_table(rows=1, cols=3)
        _header_row(table.rows[0].cells, ["面部分区", "评估状态", "主要指标"])
        for region in regions:
            cells = table.add_row().cells
            detail = _compact_metrics(region.get("指标", {}))
            vals = [region.get("分区名称", ""), region.get("评估状态", ""), detail]
            for cell, val in zip(cells, vals):
                cell.text = str(val)
                _set_cell_font(cell, size=8)
    else:
        doc.add_paragraph("本章节无可用分区结果。")

    doc.add_heading("日常管理提示", level=2)
    doc.add_paragraph(str(module.get("日常管理提示") or "本章节暂无专用日常提示。"))
    limitations = module.get("医学局限性", [])
    if limitations:
        doc.add_heading("成像与医学局限性", level=2)
        for text in limitations:
            doc.add_paragraph(str(text), style="List Bullet")


def _add_final_notes(doc: DocumentType, payload: dict[str, Any]) -> None:
    doc.add_heading("日常使用与复测说明", level=1)
    notes = [
        "本报告用于理解面部可见外观、安排日常清洁、保湿和防晒重点，并观察同条件复测趋势。",
        "复测应尽量保持相同设备、光照、距离、角度、曝光、面部表情、清洁状态及产品使用情况。",
        "底妆、遮瑕、防晒和护肤品残留，局部高光与阴影，毛发遮挡，图像压缩及分辨率不足均可能影响结果。",
        "若出现持续疼痛、灼热、明显红肿、破溃或快速变化，应咨询专业人士。",
    ]
    for note in notes:
        doc.add_paragraph(note, style="List Bullet")
    doc.add_heading("医学局限性", level=1)
    for limitation in payload.get("医学局限性", []):
        doc.add_paragraph(str(limitation), style="List Bullet")
    info = payload.get("报告信息", {})
    doc.add_paragraph(f"指标版本：{payload.get('指标版本')}    模板版本：{payload.get('模板版本')}    报告编号：{info.get('报告编号')}")


def _header_row(cells: Iterable[Any], labels: list[str]) -> None:
    for cell, label in zip(cells, labels):
        cell.text = label
        _shade(cell, CYAN)
        _set_cell_font(cell, bold=True, color="FFFFFF", size=9)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _key_value_table(doc: DocumentType, rows: list[tuple[str, Any]]) -> None:
    table = doc.add_table(rows=0, cols=2)
    for key, val in rows:
        cells = table.add_row().cells
        cells[0].text = str(key)
        cells[1].text = str(val)
        _shade(cells[0], LIGHT)
        _set_cell_font(cells[0], bold=True)
        _set_cell_font(cells[1])


def _compact_metrics(metrics: dict[str, Any], limit: int = 8) -> str:
    if not isinstance(metrics, dict) or not metrics:
        return "无"
    items = []
    for key, val in metrics.items():
        if val in ("", None):
            continue
        if isinstance(val, dict) and "value" in val:
            rendered = _format_metric_value(val)
            unit = _clean_visible(val.get("unit", ""))
            if unit and unit not in rendered:
                rendered = f"{rendered} {unit}"
        else:
            rendered = _format_value(val)
        items.append(f"{key}: {rendered}")
        if len(items) >= limit:
            break
    return "；".join(items) if items else "无"


def _format_value(value: Any) -> str:
    if isinstance(value, dict):
        return "；".join(f"{k}: {_format_value(v)}" for k, v in value.items())
    if isinstance(value, list):
        return "、".join(_format_value(v) for v in value)
    if isinstance(value, float):
        return f"{value:.6f}".rstrip("0").rstrip(".")
    return str(value)


def _shade(cell: Any, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_font(cell: Any, *, bold: bool = False, color: str = DARK, size: float = 9.5) -> None:
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.bold = bold
            run.font.name = "微软雅黑"
            run.font.size = Pt(size)
            run.font.color.rgb = RGBColor.from_string(color)
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def _finalize_tables(doc: DocumentType) -> None:
    """Apply explicit Word grid borders and readable cell spacing to every table."""
    for table in doc.tables:
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        tbl_pr = table._tbl.tblPr
        borders = tbl_pr.find(qn("w:tblBorders"))
        if borders is None:
            borders = OxmlElement("w:tblBorders")
            tbl_pr.append(borders)
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            element = borders.find(qn(f"w:{edge}"))
            if element is None:
                element = OxmlElement(f"w:{edge}")
                borders.append(element)
            element.set(qn("w:val"), "single")
            element.set(qn("w:sz"), "8")
            element.set(qn("w:space"), "0")
            element.set(qn("w:color"), "8CCFDC" if edge.startswith("inside") else "41C6DE")

        for row_index, row in enumerate(table.rows):
            tr_pr = row._tr.get_or_add_trPr()
            cant_split = tr_pr.find(qn("w:cantSplit"))
            if cant_split is None:
                tr_pr.append(OxmlElement("w:cantSplit"))
            if row_index == 0:
                header = tr_pr.find(qn("w:tblHeader"))
                if header is None:
                    tr_pr.append(OxmlElement("w:tblHeader"))
            elif row_index % 2 == 0:
                for cell in row.cells:
                    _shade(cell, "F5FBFC")
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                _set_cell_margins(cell, top=80, start=100, bottom=80, end=100)


def _set_cell_margins(cell: Any, *, top: int, start: int, bottom: int, end: int) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, amount in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(amount))
        node.set(qn("w:type"), "dxa")
