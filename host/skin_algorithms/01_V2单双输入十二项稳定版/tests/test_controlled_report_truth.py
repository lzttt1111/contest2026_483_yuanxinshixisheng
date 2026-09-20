from __future__ import annotations

from docx import Document

from src.aisia_medical_report.report import _add_doctor_module
from src.aisia_medical_report.controlled_report_truth import (
    apply_complete_scoring,
    apply_oil_metric_layout,
)
def test_complete_truth_replaces_frozen_oil_porphyrin_and_scores() -> None:
    payload = {
        "报告信息": {},
        "检测模块": [
            {"模块编号": f"{index:02d}", "核心指标": [], "医生结果分组": [], "分区指标": []}
            for index in range(1, 12)
        ],
    }
    payload["检测模块"][1].update({
        "结果摘要": "旧紫质1616",
        "核心指标": [{"name": "毛囊荧光目标数量", "value": 1616}],
    })
    payload["检测模块"][0].update({
        "综合得分": 40.48,
        "程度等级": "中度",
        "正式评分说明": {
            "summary_text": "本次评分由当前检测结果重算。",
            "drivers": [],
        },
    })
    complete = {
        "detector_results": {
            "surface_gloss": {"metrics": {
                "doctor_core_inputs": {
                    "surface_gloss_coverage": {"gloss_area_ratio": 0.0387, "high_gloss_area_ratio": 0.0013},
                    "surface_gloss_intensity": {"p50_gloss_intensity": 0.46, "p90_gloss_intensity": 0.62},
                    "surface_gloss_continuity": {"largest_gloss_component_area_ratio": 0.013},
                },
                "full_face": {"patch_count": 24},
            }},
            "porphyrin": {"metrics": {"porphyrin": {"总体指标": {
                "核心指标": {
                    "范围与负担": {"单位面积密度（个/10万有效皮肤像素）": 192.83, "特征面积占比": 0.0289},
                    "信号强度": {"实例P50强度（0～1）": 0.8294, "实例P90强度（0～1）": 0.8533},
                },
                "辅助指标": {"数量与密度": {"特征数量（个）": 371}},
            }}}},
        },
        "scoring_results": {
            "scoring_profile_version": "production_proxy_v1",
            "calibration_boundary": {"statement": "三锚点工程代理分"},
            "module_scores": {
                "oil_tendency": {"score": 73.25, "grade": "较明显", "score_valid": True, "groups": {}},
            },
        },
    }

    apply_oil_metric_layout(payload, complete)
    apply_complete_scoring(payload, complete)

    oil = payload["检测模块"][1]
    assert "371" in oil["结果摘要"]
    assert any(row["name"] == "紫质目标数量" and row["value"] == 371 for row in oil["核心指标"])
    assert oil["综合得分"] is None
    assert oil["完整评分值"] is None
    assert oil["score_valid"] is False
    assert oil["Word评分展示"] is False
    pores = payload["检测模块"][0]
    assert pores["综合得分"] == 40.48
    assert pores["程度等级"] == "中度"
    assert pores["正式评分说明"]["summary_text"] == "本次评分由当前检测结果重算。"
    assert "production_proxy_v1" not in str(payload)


def test_zero_gloss_region_keeps_count_zero_but_hides_p90() -> None:
    payload = {
        "报告信息": {},
        "检测模块": [
            {
                "模块编号": f"{index:02d}",
                "核心指标": [],
                "医生结果分组": [],
                "分区指标": [],
            }
            for index in range(1, 12)
        ],
    }
    complete = {
        "detector_results": {
            "surface_gloss": {
                "metrics": {
                    "overall_metrics": {"core_metrics": {}},
                    "region_metrics": [{
                        "analysis_region": "画面左外侧面颊",
                        "evaluation_status": "ASSESSABLE",
                        "core_metrics": {
                            "scope_and_morphology": {"gloss_area_ratio": 0.0},
                            "signal_intensity": {"p90_gloss_intensity": 0.0},
                            "count_and_density": {"gloss_patch_count": 0},
                        },
                    }],
                }
            },
            "porphyrin": {"metrics": {}},
        }
    }

    apply_oil_metric_layout(payload, complete)

    metrics = payload["检测模块"][1]["分区指标"][0]["指标"]
    assert metrics["油光区域数量"] == 0
    assert metrics["P90油光强度"] == "不可评估"

    document = Document()
    _add_doctor_module(
        document,
        {
            "模块编号": "02",
            "模块名称": "油脂分泌倾向",
            "医生评估状态": "检测结果",
            "数据来源": [],
            "结果摘要": "零目标分区语义测试",
            "结果图": [],
            "核心指标": [],
            "医生结果分组": [],
            "分区指标": [{
                "分区名称": "画面左外侧面颊",
                "状态": "检测完成",
                "指标": metrics,
            }],
        },
        2,
    )
    word_text = "\n".join(
        " | ".join(cell.text for cell in row.cells)
        for table in document.tables
        for row in table.rows
    )
    assert "P90油光强度: —" in word_text


def test_doctor_module_preserves_identical_core_and_named_group_tables() -> None:
    metric = {
        "name": "下颌线连续比例",
        "value": 0.8324,
        "unit": "0～1",
        "display_type": "ENGINEERING_0_1",
        "explanation": "表示相邻特征在有效区域内的连续程度。",
    }
    document = Document()

    _add_doctor_module(
        document,
        {
            "模块编号": "11",
            "模块名称": "面部轮廓紧致度",
            "医生评估状态": "检测结果",
            "数据来源": [],
            "结果摘要": "核心与医生分组属于不同正式章节。",
            "结果图": [],
            "核心指标": [metric],
            "医生结果分组": [{
                "title": "面部轮廓几何测量",
                "metrics": [dict(metric)],
            }],
            "分区指标": [],
        },
        11,
    )

    # 状态表、全面部核心表和具名医生分组表必须分别保留。
    assert len(document.tables) == 3
    headings = [paragraph.text for paragraph in document.paragraphs]
    assert "全面部核心指标" in headings
    assert "面部轮廓几何测量" in headings
    metric_tables = [
        table for table in document.tables
        if table.cell(0, 0).text == "核心指标"
    ]
    assert len(metric_tables) == 2
    assert [cell.text for cell in metric_tables[0].rows[1].cells] == [
        cell.text for cell in metric_tables[1].rows[1].cells
    ]
